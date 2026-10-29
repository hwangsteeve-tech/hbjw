"""한국가스기술공사 「AI 기반 실시간 가스배관망 안전관리시스템 시범구축 용역」 기타 제출서류(DOCX).

제안요청서 [서식1] 제안사 일반현황 · [서식2] 사업수행 조직 · [서식3] 계량평가 자가 점검표 · [서식4] PM 사업 수행경험,
공고문 3.가.2) 제안서 발표 대리인 위임장(대표자 불참 시), 제출서류 점검표를 한 파일로 만든다.
회사 정보는 proposal_kit/company/h9_company.yaml 기준, 확인되지 않은 값은 【 】로 남긴다.

    python proposals/2026-11_kgt-ai-pipeline-safety/build_forms.py
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parents[1]))
from proposal_kit.company import load_company, document_author  # noqa: E402

FONT = "맑은 고딕"
PROJECT = "AI 기반 실시간 가스배관망 안전관리시스템 시범구축 용역"
co = load_company()
g = co["general"]


def font(run, size=10, bold=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if "【" in run.text:
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        run.bold = True


def para(doc, text, size=10, bold=False, align=None, after=4):
    p = doc.add_paragraph()
    r = p.add_run(text)
    font(r, size, bold)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


def table(doc, rows, widths=None, head=True, size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = widths or [17 / len(rows[0])] * len(rows[0])
    tblPr = t._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(widths) / 2.54 * 1440)))
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j] / 2.54 * 1440)))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            font(r, size, bold=(head and i == 0))
            if widths:
                cell.width = Cm(widths[j])
    doc.add_paragraph()
    return t


def title(doc, form, name):
    para(doc, form, 10, True)
    para(doc, name, 15, True, WD_ALIGN_PARAGRAPH.CENTER, 10)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2))

    # 표지 · 점검표
    para(doc, PROJECT, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 6)
    para(doc, "기타 제출서류 ([서식1] ~ [서식4] · 발표 대리인 위임장)", 12, False, WD_ALIGN_PARAGRAPH.CENTER, 14)
    para(doc, "제출서류 점검표 (공고 제2026-47호 · 제안요청서 기준)", 11, True)
    table(doc, [
        ["번호", "서류", "제출 방법", "파일"],
        ["1", "정성제안서 (업체명 표기)", "나라장터 평가분야 '정성제안서' (PDF)", "kgt_정성제안서"],
        ["2", "정성제안서 (업체명 미표기)", "나라장터 평가분야 '정성제안서' (PDF)", "블라인드_정성제안서"],
        ["3", "정량제안서", "나라장터 평가분야 '정량제안서' (PDF)", "kgt_정량제안서"],
        ["4", "제안요약서 · 발표자료 (업체명 표기)", "나라장터 평가분야 '발표자료' (PDF)", "kgt_발표자료"],
        ["5", "제안요약서 · 발표자료 (업체명 미표기)", "나라장터 평가분야 '발표자료' (PDF)", "블라인드_발표자료"],
        ["6", "계량평가 자가 점검표 [서식3] · 신용평가등급 확인서", "정량제안서와 함께 별도 제출 (마감 전)", "본 문서 · 정량제안서"],
        ["7", "제안서 발표 대리인 위임장 · 재직증명서", "대표자 불참 시 (발표자 확인서류 첨부)", "본 문서"],
        ["8", "청렴계약이행서약서 · 담합입찰 방지서약서 · 정부권장사항 이행서약서", "전자입찰서 송신으로 갈음 (공고 9)", "계약특수조건_각종서약서"],
        ["9", "입찰보증금 지급각서", "전자입찰서 송신으로 갈음 (공고 7)", "-"],
    ], widths=[1.2, 7, 5.5, 3.5], size=9)
    para(doc, "※ 제출서류는 PDF 형식, 총 용량 200MB 이내 / 제안서 제출 2026.11.05(목) 10:00 ~ 11.09(월) 10:00 / 입찰참가자격 등록 마감 2026.11.06(금) 18:00", 9)
    para(doc, "※ 주민등록번호는 앞자리만 기재 / 전화 · 구두 질의 답변은 법적 효력 없음 (AX혁신실 042-600-8297)", 9)
    page_break(doc)

    # [서식1]
    title(doc, "[서식1]", "제안사 일반 현황")
    first_year, first_month = g["founded"].replace("월", "").split("년")
    table(doc, [
        ["회 사 명", g["name"], "대 표 자", g["ceo"]],
        ["사 업 분 야", g["business"], "연 락 처", f"전화 : {g['tel']}  FAX : {g['fax']}"],
        ["주 소", g["address"], "회사설립연도", f"{first_year.strip()}년 {first_month.strip()}월"],
        ["해당부문 사업기간", g["biz_period"], "사업자등록번호", g["biz_no"]],
    ], widths=[3, 5.5, 3, 5.5], head=False)
    para(doc, "<주요연혁>", 10, True)
    table(doc, [["연도", "주요 연혁"]] + [[str(y), " / ".join(items)] for y, items in co["history"]],
          widths=[2, 15], size=8.5)
    page_break(doc)

    # [서식2]
    title(doc, "[서식2]", "사업수행 조직")
    para(doc, "□ 조직도", 10, True)
    table(doc, [
        ["구분", "역할", "성명", "인원"],
        ["사업PM (사업관리책임자)", "사업 총괄 · 일정 · 위험 · 발주기관 보고, 제안 발표, 현장실증 총괄", "【성명】", "1명"],
        ["품질보증 담당", "단계별 품질 점검 · 산출물 검토 (수행팀과 분리)", "【성명】", "1명"],
        ["AI 모델팀", "탐지대상 정의 · 라벨링, 1차 · 2차 모델 학습 · 평가, 마스킹 모델", "【성명】", "【 】명"],
        ["현장 단말팀", "단말 7세트 구성 · 설치, GNSS 동기화 · 통신 · 단말 보안", "【성명】", "【 】명"],
        ["플랫폼팀", "통합관제 · GIS · 기준정보 · 계정권한 · 연계 · 외부 API, UX/UI", "【성명】", "【 】명"],
        ["품질 · 보안", "단계별 · 성능 · 부하 시험, 시큐어코딩 · 취약점 · 개인정보 점검", "【성명】", "【 】명"],
    ], widths=[3.6, 8.6, 2.6, 2.2], size=9)
    para(doc, "※ 공동수급 · 하도급 없이 당사 인력이 직접 수행 (발주기관: AX혁신실 추진총괄 · 공급기술처 · 실증대상 지사 · 정보보안센터)", 9)
    page_break(doc)

    # [서식3]
    title(doc, "[서식3]", "계량평가 자가 점검표")
    table(doc, [
        ["평가항목", "평가요소", "배점", "평점", "비고"],
        ["경영상태 (5점)", "경영상태 (신용평가 등급 확인서)", "5", "5", "기업신용평가등급 BBB0"],
        ["계", "", "5", "5", ""],
    ], widths=[3.4, 6, 1.6, 1.6, 4.4])
    para(doc, "※ 근거 자료 사본 첨부 — 신용평가등급 확인서 【신용평가기관 · 평가일 · 유효기간】", 9)
    para(doc, "※ 본 사업은 추정가격 고시금액 미만 소규모 계약으로 유사사업 수행실적은 평가항목에서 제외", 9)
    para(doc, "※ 신용평가등급(AAA ~ BBB0) 5점 / (BBB- ~ BB-) 4.75점 / (B+ ~ B-) 4.5점 / (CCC+ 이하) 3.5점", 9)
    page_break(doc)

    # [서식4]
    title(doc, "[서식4]", "사업관리책임자(PM)의 사업 수행경험")
    table(doc, [
        ["성 명", "【성명】", "소 속", f"{g['name']} 【부서】"],
        ["직 책", "【직책】", "연 령", "【 】세"],
        ["학 력", "【 】대학교 【 】전공", "해당분야근무경력", "【 】년 【 】개월"],
        ["", "【 】대학원 【 】전공", "자 격 증", "【자격증】"],
    ], widths=[3, 5.5, 3.5, 5], head=False)
    table(doc, [["사업명", "사업기간", "사업개요", "사업금액", "수행임무", "발주기관"]] +
          [["【사업명】", "【 ~ 】", "【개요】", "【 】", "【 】", "【발주기관】"] for _ in range(5)],
          widths=[3.6, 2.6, 3.8, 2, 2, 3], size=9)
    para(doc, "※ 제안서에 명시된 사업관리책임자(PM)가 제안 발표를 직접 수행 (미수행 시 서면평가)", 9)
    page_break(doc)

    # 발표 대리인 위임장
    title(doc, "[공고 3.가.2)]", "제안서 발표 대리인 위임장")
    para(doc, "※ 대표자가 제안 발표에 참석하지 않는 경우 제출 (재직증명서 1부 첨부)", 9, after=10)
    table(doc, [
        ["입찰건명", PROJECT],
        ["공고번호", "한국가스기술공사 공고 제2026-47호 (R26BK01796859-000)"],
        ["발표 일시 · 장소", "2026. 11. 10. 【시간 — 개별 공지】 / 대전광역시 유성구 대덕대로 1227 한국가스기술공사 본사"],
    ], widths=[4, 13], head=False)
    para(doc, "□ 대리인 (발표자)", 10, True)
    table(doc, [
        ["성 명", "【성명】", "생년월일", "【 】 (앞자리만)"],
        ["소 속 · 직위", f"{g['name']} 【부서 · 직위】", "연락처", "【 】"],
        ["역 할", "사업관리책임자(PM)", "첨부", "재직증명서 1부"],
    ], widths=[3, 5.5, 3, 5.5], head=False)
    para(doc, "위 사람을 본 입찰의 제안서 발표 대리인으로 정하여 제안서 설명 및 질의응답에 관한 일체의 권한을 위임합니다.", 10, after=16)
    para(doc, "2026년    11월    일", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 16)
    table(doc, [
        ["위임인", f"상호 : {g['name']}"],
        ["", f"주소 : {g['address']}"],
        ["", f"대표자 : {g['ceo']}   (인)"],
    ], widths=[3, 14], head=False)
    para(doc, "한국가스기술공사 사장 귀하", 12, True, WD_ALIGN_PARAGRAPH.CENTER)

    cp = doc.core_properties
    cp.author = cp.last_modified_by = document_author()
    cp.title = f"{PROJECT} 기타 제출서류"
    out = HERE / "기타제출서류_서식1-4_위임장.docx"
    doc.save(str(out))
    print("저장:", out)


if __name__ == "__main__":
    main()
