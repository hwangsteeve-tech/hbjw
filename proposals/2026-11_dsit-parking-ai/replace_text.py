"""사용자 확정본 PPT(dsit_proposal.pptx)의 글자만 바꾼다 — 레이아웃 · 글꼴 · 크기는 그대로 둔다.

사용: python proposals/2026-11_dsit-parking-ai/replace_text.py "바꿀 문구" "새 문구" [--slide 9]
- 문구가 한 run 안에 있으면 그 run만, 여러 run에 걸치면 첫 run에 합쳐 넣고 나머지 run을 비운다
  (첫 run의 서식을 따름). 표 · 그룹 도형 안의 글자도 찾는다.
- 바뀐 곳이 없으면 오류로 끝난다.
"""
import argparse
import sys
from pathlib import Path

from pptx import Presentation

PPT = Path(__file__).with_name("dsit_proposal.pptx")


def _frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # 그룹
            yield from _frames(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def replace(old, new, slide_no=None):
    prs = Presentation(PPT)
    hits = []
    for i, slide in enumerate(prs.slides, 1):
        if slide_no and i != slide_no:
            continue
        for tf in _frames(slide.shapes):
            for p in tf.paragraphs:
                runs = p.runs
                text = "".join(r.text for r in runs)
                if old not in text:
                    continue
                for r in runs:  # 한 run 안에 있는 경우
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        break
                else:  # 여러 run에 걸친 경우
                    runs[0].text = text.replace(old, new)
                    for r in runs[1:]:
                        r.text = ""
                hits.append(i)
    if not hits:
        sys.exit(f"찾지 못함: {old!r}")
    prs.save(PPT)
    print(f"{len(hits)}곳 변경 (쪽: {sorted(set(hits))})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("old")
    ap.add_argument("new")
    ap.add_argument("--slide", type=int)
    a = ap.parse_args()
    replace(a.old, a.new, a.slide)
