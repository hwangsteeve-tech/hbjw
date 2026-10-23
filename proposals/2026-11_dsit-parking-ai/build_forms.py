"""동성아이텍 <별첨 2> 제출서류 양식 1~6 작성 (Word).

원본: AI 기반 주차 수요예측·자동감면 서비스 개발 용역_제출양식_261019.hwp
내용 기준: content.yaml(정성제안서) Ⅱ-2 · Ⅱ-3 · Ⅷ-1, proposal_kit/company/h9_company.yaml
【 】 = 확인 후 기입 (붉은색 굵은 글씨)

실행: python proposals/2026-11_dsit-parking-ai/build_forms.py
"""
import re
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
ROOT = HERE.parents[1]
COMPANY = yaml.safe_load((ROOT / "proposal_kit/company/h9_company.yaml").read_text(encoding="utf-8"))
FONT = "맑은 고딕"
PROJECT = "AI 기반 주차 수요예측·자동감면 서비스 개발 용역"
CLIENT = "주식회사 동성아이텍"
DATE = "2026년  11월  25일"
ON, OFF = "■", "□"

# [양식 1] 유사사업 수행실적 — 정성제안서 Ⅱ-3과 같은 목록. 회사 실적 목록의 사업명 · 기간 · 금액(원)을 그대로 옮기고, 양식에는 백만원 단위로 적음
RECORDS = [
    ("AI개발", "서울아산병원 AI-IDP (통합데이터플랫폼) 구축사업", "에스케이㈜", "2026.06.10~2027.11.27", "883,225,000", "수행중", "계약서 첨부"),
    ("AI개발", "SK가스26년Trading AI 사업(PP)", "에스케이㈜", "2026.07.01~2026.07.24", "32,324,620", "완료", "실적증명서 첨부"),
    ("AI개발", "SKHy '26년 용인 Y1(M17) 제조 DT 시스템 Ph1_FAB MOS_원가선투입", "에스케이㈜", "2026.05.01~2026.06.30", "34,551,600", "완료", "실적증명서 첨부"),
    ("AI개발", "Trading AI Agent 개발 및 ML/DL 분석_선투입", "에스케이㈜", "2026.04.08~2026.06.30", "129,385,380", "완료", "실적증명서 첨부"),
    ("AI개발", "SKHy '26년 M14 차세대 제조 시스템 확산_FABMOS_원가선투입", "에스케이㈜", "2026.03.16~2026.06.30", "32,853,200", "완료", "실적증명서 첨부"),
    ("AI개발", "SKHy '26년 M15X Ph4 차세대 제조 시스템 _FAB MOS", "에스케이㈜", "2026.03.01~2026.06.30", "27,440,000", "완료", "실적증명서 첨부"),
    ("공공", "25년 국립중앙박물관 모두 플랫폼 개선 및 유지보수", "국립중앙박물관", "2025.07.01~2025.12.31", "86,200,000", "완료", "실적증명서 첨부"),
    ("공공", "2024년 국립중앙박물관 교육 플랫폼 [모두] 유지보수 사업", "국립중앙박물관", "2024.07.18~2025.01.14", "100,000,000", "완료", "실적증명서 첨부"),
]

# [양식 3] 투입인력 총괄표 — 정성제안서 Ⅷ-1과 같은 구성
STAFF = [
    ("PM", "AI 서비스 · 공공 플랫폼 사업관리", "M1~M6", "3.0"),
    ("ML 엔지니어 (책임)", "수요 · 점유 · 복귀시간 예측 모델", "M1~M6", "5.0"),
    ("ML 엔지니어", "자동 공급 · 동적 요금 최적화", "M2~M5", "3.0"),
    ("LLM · NLP 엔지니어", "조례 파싱 · 감면 Rule 추출", "M1~M6", "4.0"),
    ("백엔드 · Rule 엔진", "Rule DB · 감면 판단 · 정산 · API 연계", "M2~M6", "4.0"),
    ("데이터 엔지니어", "데이터 파이프라인 · MLOps", "M1~M5", "2.5"),
    ("QA · 보안", "테스트 · 개인정보 보호", "M4~M6", "1.0"),
]


# ---------------------------------------------------------------- 서식 도우미

def _font(run, size, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_runs(par, text, size=10, bold=False, color=None):
    """【 】 칸은 붉은색 굵은 글씨로 넣는다."""
    for part in re.split(r"(【[^】]*】)", text):
        if not part:
            continue
        blank = part.startswith("【")
        _font(par.add_run(part), size, bold or blank, "C00000" if blank else color)


def para(doc, text="", size=10, bold=False, align=None, before=0, after=4, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.15
    if align:
        p.alignment = align
    add_runs(p, text, size, bold, color)
    return p


def _shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(doc, header, rows, widths, size=9, aligns=None, row_h=0.8):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # Word는 칸 너비, 한글 · LibreOffice는 표 격자(gridCol) 너비를 따르므로 둘 다 맞춘다
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for r, values in enumerate([header] + rows):
        tr = t.rows[r]
        tr.height = Cm(row_h)
        for c, v in enumerate(values):
            cell = tr.cells[c]
            cell.width = Cm(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r == 0 or not aligns or aligns[c] == "c") else WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, str(v), size, bold=(r == 0))
            if r == 0:
                _shade(cell, "D9D9D9")
    para(doc, after=2)
    return t


def note(doc, text):
    para(doc, text, 9, after=2)


def form_title(doc, no, title, first=False):
    if not first:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, f"[양식 {no}]", 10, bold=True, after=6)
    para(doc, title, 18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=14)


def confirm(doc, text="상기 기재내용은 사실과 다름이 없음을 확인합니다."):
    para(doc, text, 11, align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=14)
    para(doc, DATE, 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    para(doc, f"업체명 : {COMPANY['general']['name']}", 11, align=WD_ALIGN_PARAGRAPH.RIGHT, after=4)
    para(doc, f"대표자 : {COMPANY['general']['ceo']}          (직인)", 11, align=WD_ALIGN_PARAGRAPH.RIGHT, after=24)
    para(doc, f"{CLIENT} 귀하", 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def mil(won):
    """원 금액을 양식 단위(백만원)로 바꾼다. 반올림 없이 정확한 값 (예: 32,324,620 → 32.32462)."""
    from decimal import Decimal
    v = Decimal(won.replace(",", "")) / Decimal(1_000_000)
    whole, _, frac = f"{v:f}".partition(".")
    frac = frac.rstrip("0")
    return f"{int(whole):,}" + (f".{frac}" if frac else "")


def check(yes):
    return f"{ON} 해당  {OFF} 미해당" if yes is True else f"{OFF} 해당  {ON} 미해당" if yes is False else f"{OFF} 해당  {OFF} 미해당"


# ---------------------------------------------------------------- 양식

def build(out):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

    # 제출서류 목록
    para(doc, "<별첨 2> 제출서류 양식", 10, bold=True, after=6)
    para(doc, "제출 서류 양식 및 증빙자료", 18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=14)
    table(doc, ["No", "제출 서류", "양식", "부수", "첨부 증빙", "준비"], [
        ["1", "유사사업 수행실적", "양식 1", "1부", "계약서 · 실적증명서 (8건)", "【 】"],
        ["2", "연구개발 조직 현황", "양식 2", "1부", "기업부설연구소 인정서", "【 】"],
        ["3", "투입인력 총괄표", "양식 3", "1부", "4대보험 가입내역 확인서 또는 재직증명서 (7명)", "【 】"],
        ["4", "정책적 우대기업 및 기술역량 현황", "양식 4", "1부", "해당 없음 (5개 항목 미해당)", "-"],
        ["5", "서약서", "양식 5", "1부", "-", "【 】"],
        ["6", "개인정보 수집 · 이용 동의서", "양식 6", "7부", "투입인력별 서명", "【 】"],
        ["7", "사업자등록증 사본", "-", "1부", "-", "【 】"],
        ["8", "기타증빙 서류", "-", "1식", "-", "【 】"],
    ], [1.0, 4.6, 1.6, 1.2, 6.0, 1.6], aligns="cllclc", row_h=0.9)
    note(doc, "※ 증빙자료를 첨부하지 않은 항목은 인정하지 않음 (양식 1 · 2 · 4), 투입인력은 실제 투입 예정 인력 기준 (양식 3).")
    note(doc, "※ 날짜 · 직인 · 서명과 붉은색 【 】 칸은 제출 전에 채움.")

    # [양식 1]
    form_title(doc, 1, "유사사업 수행실적")
    para(doc, "(단위: 백만원)", 9, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2)
    table(doc, ["구분", "사업명", "발주처", "사업기간", "계약금액", "상태", "비고"],
          [[*r[:3], r[3].replace("~", "~\n"), mil(r[4]), *r[5:]] for r in RECORDS], [1.8, 3.9, 1.8, 2.6, 2.6, 1.4, 2.6], size=8.5,
          aligns="clcccc" + "c", row_h=0.85)
    note(doc, "※ 구분은 국가R&D, 지자체 실증, 공공분야 정책연구 등 해당 분야를 기재함.")
    note(doc, "※ 기재한 수행실적은 계약서, 실적증명서, 협약서, 결과통보서 등 객관적 증빙자료를 첨부하여야 하며, 미첨부 시 해당 실적은 가점 산정에서 제외함.")
    confirm(doc)

    # [양식 2]
    form_title(doc, 2, "연구개발 조직 현황")
    table(doc, ["구분", "해당 여부", "인정·등록번호", "인정·등록일", "발급기관"], [
        ["연구개발전담부서", check(False), "-", "-", "-"],
        ["기업부설연구소", check(True), "제2025112883호", "2025.09.10", "한국산업기술진흥협회"],
        ["전문연구사업자", check(False), "-", "-", "-"],
    ], [3.4, 3.4, 3.6, 2.8, 3.8], row_h=1.2)
    note(doc, "※ 해당 연구개발 조직의 인정서, 신고증 또는 등록증 등 객관적 증빙자료를 첨부하여야 하며, 미첨부 시 해당 항목은 인정하지 않음.")
    note(doc, "※ 복수 항목에 해당하는 경우 가장 높은 배점 1개만 적용함.")
    confirm(doc)

    # [양식 3]
    form_title(doc, 3, "투입인력 총괄표")
    rows = [[role, "【성명】", "【직위】", task, "【학력】", "【 】년", period, f"{mm} M/M"]
            for role, task, period, mm in STAFF]
    rows.append(["합계", f"{len(STAFF)}명", "-", "-", "-", "-", "M1~M6", f"{sum(float(s[3]) for s in STAFF):.1f} M/M"])
    table(doc, ["구분", "이름", "직위", "담당업무", "학력", "경력", "투입기간", "비고"], rows,
          [3.2, 1.5, 1.4, 4.4, 1.6, 1.3, 1.6, 1.9], size=8.5, aligns="ccclcccc", row_h=0.95)
    note(doc, "※ 본 사업에 실제 투입 예정인 인력을 기준으로 작성하며, 재직 여부를 확인할 수 있는 증빙자료를 반드시 첨부하여야 함.")
    confirm(doc)

    # [양식 4] — 경영기획팀 회신(2026-10-19): 5개 항목 모두 미해당
    form_title(doc, 4, "정책적 우대기업 및 기술역량 현황")
    table(doc, ["구분", "해당 여부", "증빙자료", "비고"], [
        ["벤처기업 등 기술기업 인증", check(False), "-", "-"],
        ["여성기업", check(False), "-", "-"],
        ["청년대표기업", check(False), "-", "-"],
        ["사회적경제기업", check(False), "-", "-"],
        ["지역 사업수행 기반", check(False), "-", "-"],
    ], [5.0, 3.6, 4.2, 4.2], row_h=1.1)
    note(doc, "※ 해당 항목의 관련 증빙자료를 반드시 첨부하여야 하며, 미첨부 시 해당 항목은 인정하지 않음.")
    confirm(doc)

    # [양식 5]
    form_title(doc, 5, "서 약 서")
    para(doc, f"『{PROJECT}』입찰과 관련하여 제출하는 서류 및 제 증빙자료는 신의․성실의 원칙에 의하여 작성 제출합니다.",
         12, before=24, after=14)
    para(doc, "만일 제출한 서류 및 증빙자료가 부정한 방법으로 작성되었거나 허위 기재 사실이 확인될 경우 참가 자격에서 제외하여도 이의를 제기하지 않을 것을 서약합니다.",
         12, after=60)
    confirm(doc, "")

    # [양식 6] — 참여인력별로 각각 작성
    for role, *_ in STAFF:
        form_title(doc, 6, "개인정보 수집∙이용 동의서")
        para(doc, f"{CLIENT}은 『{PROJECT}』과 관련하여 「개인정보 보호법」에 따라 아래와 같이 개인정보를 수집·이용하고자 합니다.", 11, after=12)
        for head, body in [("수집·이용 목적", ["제안서 평가 및 참여인력의 학력·경력 등 자격사항 확인"]),
                           ("수집·이용 항목", ["성명, 연령, 학력·전공, 관련경력, 자격사항, 담당업무"]),
                           ("보유·이용기간", ["수집·이용 목적 달성 후 지체 없이 파기함. 다만, 관련 법령에 따라 보존이 필요한 경우에는 해당 기간 동안 보관함."])]:
            para(doc, f"□ {head}", 11, bold=True, before=6, after=2)
            for b in body:
                para(doc, f"  ㅇ {b}", 11, after=4)
        para(doc, "※ 개인정보 수집·이용에 대한 동의를 거부할 권리가 있습니다.", 10, before=12, after=2)
        para(doc, "다만, 동의를 거부할 경우 참여인력 평가에 제한이 있을 수 있습니다.", 10, after=14)
        para(doc, f"위와 같이 귀하의 개인정보를 수집·이용하는 것에 동의하십니까?   {OFF} 동의함      {OFF} 동의하지 않음",
             11, bold=True, after=30)
        para(doc, DATE, 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
        para(doc, f"동 의 자 : 【{role} 성명】          (서명 또는 인)", 11, align=WD_ALIGN_PARAGRAPH.RIGHT, after=24)
        note(doc, "※ 본 양식은 참여인력별로 각각 작성함.")

    author = COMPANY["document_author"]
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = author
    doc.core_properties.title = f"{PROJECT} 제출서류 양식"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build(HERE / "제출양식_동성아이텍.docx")
