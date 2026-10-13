"""H9 제안 발표자료(16:9) PPT 생성기.

YAML 내용 파일을 읽어 표지 · 평가항목 조견표 · CONTENTS · 장 간지 · 모듈 간지 · 핵심 비주얼 · 본문 · 마무리
슬라이드를 만든다. 형식 기준은 reference/의 발표자료 예시(AI학습용데이터 통합 제공 시스템).

    python -m proposal_kit.build_presentation proposals/xxx/presentation.yaml -o 발표자료.pptx --report

본문은 headline 아래 영역을 columns(가로 비율)로 나누고, 각 열에 블록을 위에서 아래로 쌓는다.
블록 종류는 README.md 참고.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

from proposal_kit.company import set_pptx_author
from proposal_kit.style import finalize

HERE = Path(__file__).resolve().parent
BG_DIR = HERE / "templates" / "presentation"

FB = "KoPubWorld돋움체 Bold"
FM = "KoPubWorld돋움체 Medium"

NAVY = "1E3A8A"
BLUE = "1F5FD1"
SKY = "2F80ED"
DARK = "1F2A44"
GRAY_BAND = "F2F3F5"
LINE = "C9D3E3"
LIGHT = "EEF3FB"
RED = "D93025"

W, H = 33.867, 19.05
X0, X1 = 0.9, 32.97
CONTENT_TOP = 4.35
CONTENT_BOTTOM = 18.0
CONCL_H = 1.15

_MARK = re.compile(r"(\*\*.+?\*\*|!!.+?!!)")


# ---------------------------------------------------------------- 공통 유틸

def rgb(h):
    return RGBColor.from_string(h)


def runs(text):
    """'**파랑 강조**', '!!빨강 강조!!' 표기를 나눈다."""
    out = []
    for part in _MARK.split(str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], "blue"))
        elif part.startswith("!!") and part.endswith("!!"):
            out.append((part[2:-2], "red"))
        else:
            out.append((part, None))
    return out


SCALE = 1.15  # 본문(12.5pt 이하) 글자 확대 비율


def scaled(size):
    return size * SCALE if size <= 12.5 else size


def font(run, size, face=FM, color=DARK, bold=False):
    run.font.size = Pt(scaled(size))
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)


def add_text(par, text, size, face=FM, color=DARK, em_blue=BLUE, em_red=RED):
    for chunk, em in runs(text):
        r = par.add_run()
        r.text = chunk
        if em == "blue":
            font(r, size, FB, em_blue)
        elif em == "red":
            font(r, size, FB, em_red)
        else:
            font(r, size, face, color)


def pfmt(par, align="l", line=1.15, before=0, after=0):
    par.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    par.line_spacing = line
    par.space_before = Pt(before)
    par.space_after = Pt(after)


def bullet(par, char="•", mar=0.32, ind=0.28, face="Arial"):
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(int(Cm(mar))))
    pPr.set("indent", str(-int(Cm(ind))))
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", face)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def box(slide, kind, x, y, w, h, fill=None, line=None, lw=0.75, adj=None):
    shp = slide.shapes.add_shape(getattr(MSO_SHAPE, kind), Cm(x), Cm(y), Cm(w), Cm(h))
    if adj is not None:
        shp.adjustments[0] = adj
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    spPr = shp._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))
    shp.text_frame.word_wrap = True
    return shp


def fill_text(shp, lines, size, face=FB, color=DARK, align="c", anchor="m", inset=0.12, line=1.1,
              em_blue=BLUE, em_red=RED):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for s in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{s}", Cm(inset))
    lines = lines if isinstance(lines, list) else str(lines).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pfmt(p, align, line)
        add_text(p, ln, size, face, color, em_blue, em_red)
    return shp


def tbox(slide, x, y, w, h, anchor="t", inset=0.05):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for s in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{s}", Cm(inset))
    return tb


def gradient(shp, c1, c2, angle=0):
    spPr = shp._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    g = etree.Element(qn("a:gradFill"))
    g.set("rotWithShape", "1")
    gs = etree.SubElement(g, qn("a:gsLst"))
    for pos, c in ((0, c1), (100000, c2)):
        s = etree.SubElement(gs, qn("a:gs"))
        s.set("pos", str(pos))
        etree.SubElement(s, qn("a:srgbClr")).set("val", c)
    lin = etree.SubElement(g, qn("a:lin"))
    lin.set("ang", str(angle * 60000))
    lin.set("scaled", "0")
    geom = spPr.find(qn("a:prstGeom"))
    geom.addnext(g)


def alpha(shp, pct):
    sf = shp._element.spPr.find(qn("a:solidFill"))
    clr = sf[0]
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(pct * 1000)))


def est_lines(text, size, width_cm, indent=0.0):
    plain = re.sub(r"\*\*|!!", "", str(text))
    char = scaled(size) * 0.0353 * 1.0
    cap = max(1, int((width_cm - indent) / char))
    total = 0
    for seg in plain.split("\n"):
        units = sum(1.0 if ord(c) > 0x2000 else 0.55 for c in seg)
        total += max(1, -(-int(units) // cap))
    return total


def lh(size, spacing=1.2):
    return scaled(size) * 0.0353 * spacing


# ---------------------------------------------------------------- 고정 슬라이드

class Deck:
    def __init__(self, data):
        self.data = data
        self.prs = Presentation()
        self.prs.slide_width = Cm(W)
        self.prs.slide_height = Cm(H)
        self.blank = self.prs.slide_layouts[6]
        self.page_of = {}      # 평가항목 -> 쪽 목록
        self.fill = []

    def new(self):
        return self.prs.slides.add_slide(self.blank)

    def page_no(self, s):
        n = len(self.prs.slides)
        tb = tbox(s, W - 2.2, H - 0.85, 1.6, 0.6, anchor="m")
        p = tb.text_frame.paragraphs[0]
        pfmt(p, "r")
        add_text(p, str(n), 9, FM, "7F7F7F")

    # 표지
    def cover(self, m):
        s = self.new()
        tb = tbox(s, 1.8, 5.1, 25, 1.0)
        add_text(tb.text_frame.paragraphs[0], m.get("doc_type", "제안 발표 자료"), 13, FM, "8C8C8C")
        tb = tbox(s, 1.8, 6.0, 30, 2.0, anchor="m")
        add_text(tb.text_frame.paragraphs[0], m["title"], 30, FB, DARK)
        segs = [(0.0, 0.62, SKY), (0.62, 0.66, "0B1A4A"), (0.66, 1.0, "F6C77B")]
        for a, b, c in segs:
            box(s, "RECTANGLE", 2.0 + a * 21.5, 8.25, (b - a) * 21.5, 0.12, fill=c)
        tb = tbox(s, 2.4, 8.6, 20, 1.0, anchor="m")
        p = tb.text_frame.paragraphs[0]
        add_text(p, m.get("pm_label", "사업책임자 "), 12, FM, DARK)
        add_text(p, m.get("pm_name", "【성명】"), 16, FB, DARK)
        add_text(p, " " + m.get("pm_title", ""), 12, FM, DARK)
        box(s, "RECTANGLE", 2.0, 9.75, 21.5, 0.03, fill="BFBFBF")
        tb = tbox(s, 2.4, 10.3, 22, 3.0)
        for i, ln in enumerate(m.get("org_lines", [])):
            p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
            pfmt(p, after=6)
            bullet(p, "▪", 0.5, 0.45)
            add_text(p, ln, 12, FM, DARK)
        logos = m.get("logos", [])
        for i, lg in enumerate(logos):
            tb = tbox(s, 6.0 + i * 5.2, 14.4, 5.0, 1.2, anchor="m")
            p = tb.text_frame.paragraphs[0]
            pfmt(p, "c")
            add_text(p, lg, 20, FB, "111111")
        if m.get("client"):
            tb = tbox(s, 6.0, 13.0, 15, 1.0, anchor="m")
            p = tb.text_frame.paragraphs[0]
            pfmt(p, "c")
            add_text(p, m["client"], 14, FB, NAVY)
        return s

    # 평가항목 조견표 (쪽수는 빌드 후 채움)
    def eval_table(self, items):
        s = self.new()
        box(s, "RECTANGLE", 1.0, 0.75, 0.12, 0.6, fill=NAVY)
        tb = tbox(s, 1.25, 0.6, 15, 0.9, anchor="m")
        add_text(tb.text_frame.paragraphs[0], "평가항목 조견표", 14, FB, DARK)
        self._eval_slide = s
        self._eval_items = items
        return s

    def fill_eval_table(self):
        s, items = self._eval_slide, self._eval_items
        rows = [["구분", "평가항목", "평가기준", "배점", "해당 페이지"]]
        for it in items:
            pages = sorted(set(self.page_of.get(it["item"], [])))
            rows.append([it["group"], it["item"], it["criteria"], str(it["score"]),
                         it.get("page_text") or _ranges(pages)])
        rows.append(["", "", "합계", str(sum(i["score"] for i in items if isinstance(i["score"], (int, float)))), ""])
        n = len(rows)
        row_h = min(0.9, (H - 2.6) / n)
        table(s, 1.0, 1.7, W - 2.0, rows, [3.4, 5.2, 16.2, 2.0, 5.07], size=8.5 if n > 16 else 9.5,
              heights=[row_h] * n, merge_first=True, header_fill=BLUE)

    def contents(self, chapters):
        s = self.new()
        tb = tbox(s, 1.6, 0.9, 20, 2.6, anchor="m")
        add_text(tb.text_frame.paragraphs[0], "CONTENTS", 40, "Arial Black", "0B3AA4")
        n = len(chapters)
        cols = 3 if n > 4 else n
        bw, bh, gx, gy = 9.4, 5.0, 1.0, 1.35
        x0 = (W - (cols * bw + (cols - 1) * gx)) / 2
        for i, ch in enumerate(chapters):
            x = x0 + (i % cols) * (bw + gx)
            y = 5.6 + (i // cols) * (bh + gy)
            box(s, "RECTANGLE", x, y, bw, bh, line=NAVY, lw=1.25)
            cover = box(s, "RECTANGLE", x - 0.1, y - 0.6, 2.6, 1.2, fill="FFFFFF")
            tb = tbox(s, x - 0.15, y - 0.95, 1.6, 1.6, anchor="m")
            add_text(tb.text_frame.paragraphs[0], ch["num"], 30, "Times New Roman", NAVY)
            tb = tbox(s, x + 1.05, y - 0.25, 2.5, 0.6, anchor="m")
            add_text(tb.text_frame.paragraphs[0], "chapter", 9, FM, "595959")
            tb = tbox(s, x, y + 0.5, bw, bh - 0.5, anchor="m")
            p = tb.text_frame.paragraphs[0]
            pfmt(p, "c")
            add_text(p, ch["title"], 20, FB, DARK)
        return s

    def divider(self, ch):
        s = self.new()
        s.shapes.add_picture(str(BG_DIR / "divider_dark.jpg"), 0, 0, Cm(W), Cm(H))
        panel = box(s, "RECTANGLE", -0.2, 4.4, 15.8, 9.6, fill="0A1530", line="FFFFFF", lw=2.0)
        alpha(panel, 55)
        tb = tbox(s, 1.3, 5.0, 2.4, 2.6, anchor="m")
        add_text(tb.text_frame.paragraphs[0], ch["num"], 44, "Times New Roman", "FFFFFF")
        tb = tbox(s, 3.6, 5.2, 12, 2.2, anchor="m")
        add_text(tb.text_frame.paragraphs[0], ch["title"], 30, FB, "FFFFFF")
        tb = tbox(s, 2.7, 8.0, 12.5, 5.5)
        for i, t in enumerate(ch.get("toc", [])):
            p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
            pfmt(p, after=9)
            add_text(p, f"{i + 1}. {t}", 17, FM, "FFFFFF")
        return s

    def module(self, title):
        s = self.new()
        s.shapes.add_picture(str(BG_DIR / "visual_light.jpg"), 0, 0, Cm(W), Cm(H))
        ov = box(s, "RECTANGLE", 0, 0, W, H, fill="0B1A3A")
        alpha(ov, 45)
        band = box(s, "RECTANGLE", 0, 7.6, W, 3.0, fill="0B1A3A")
        alpha(band, 55)
        tb = tbox(s, 1.8, 7.6, 30, 3.0, anchor="m")
        add_text(tb.text_frame.paragraphs[0], title, 28, FM, "FFFFFF")
        return s

    def visual(self, spec):
        """핵심 비주얼: 밝은 배경 + 큰 제목 + 개요 상자."""
        s = self.new()
        s.shapes.add_picture(str(BG_DIR / "visual_light.jpg"), 0, 0, Cm(W), Cm(H))
        tb = tbox(s, 1.5, 3.0, 31, 2.3, anchor="m")
        p = tb.text_frame.paragraphs[0]
        pfmt(p, "c")
        add_text(p, spec["title"], 34, FB, "0B3AA4", em_blue="0B3AA4")
        if spec.get("subtitle"):
            tb = tbox(s, 1.5, 5.4, 31, 1.2, anchor="m")
            p = tb.text_frame.paragraphs[0]
            pfmt(p, "c")
            add_text(p, spec["subtitle"], 15, FB, DARK)
        facts = spec.get("facts", [])
        if facts:
            n = len(facts)
            bw = (W - 3.0 - 0.4 * (n - 1)) / n
            for i, f in enumerate(facts):
                x = 1.5 + i * (bw + 0.4)
                hd = box(s, "RECTANGLE", x, 12.6, bw, 1.0, fill=NAVY)
                fill_text(hd, f["label"], 11, FB, "FFFFFF")
                bd = box(s, "RECTANGLE", x, 13.6, bw, 2.4, fill="FFFFFF", line=NAVY)
                fill_text(bd, f["value"], 13 if len(str(f["value"])) < 30 else 11, FB, DARK)
        self.page_no(s)
        return s

    def closing(self, spec):
        s = self.new()
        s.shapes.add_picture(str(BG_DIR / "visual_light.jpg"), 0, 0, Cm(W), Cm(H))
        tb = tbox(s, 1.5, 6.2, 31, 2.4, anchor="m")
        p = tb.text_frame.paragraphs[0]
        pfmt(p, "c")
        add_text(p, spec.get("title", "감사합니다"), 40, FB, "0B3AA4")
        if spec.get("message"):
            tb = tbox(s, 1.5, 8.8, 31, 2.0, anchor="m")
            p = tb.text_frame.paragraphs[0]
            pfmt(p, "c")
            add_text(p, spec["message"], 16, FB, DARK)
        return s

    # ------------------------------------------------------------ 본문
    def body(self, chapter, spec):
        s = self.new()
        # 머리띠
        box(s, "RECTANGLE", 0, 0, W, 1.6, fill=GRAY_BAND)
        box(s, "PARALLELOGRAM", -0.9, 0, 2.4, 1.6, fill=SKY, adj=0.45)
        box(s, "PARALLELOGRAM", 1.05, 0, 0.5, 1.6, fill=NAVY, adj=0.9)
        tb = tbox(s, 1.75, 0.15, 24, 1.3, anchor="m")
        add_text(tb.text_frame.paragraphs[0], spec["title"], 17, FB, DARK)
        tb = tbox(s, W - 9.5, 0.2, 8.9, 1.2, anchor="m")
        p = tb.text_frame.paragraphs[0]
        pfmt(p, "r")
        add_text(p, f"{chapter['num']}. {chapter['title']}", 10, FM, DARK)
        if spec.get("req"):
            long_req = len(spec["req"]) > 44
            rb = box(s, "ROUNDED_RECTANGLE", W - 11.6, 1.8, 10.7, 1.0 if long_req else 0.7, fill="DCE6F6", adj=0.2)
            fill_text(rb, spec["req"], 8 if long_req else 9, FB, NAVY, inset=0.05)
        y = 2.0
        if spec.get("num"):
            box(s, "RECTANGLE", X0, y - 0.05, 1.0, 0.04, fill=DARK)
            tb = tbox(s, X0 - 0.05, y, 1.6, 1.15, anchor="m")
            add_text(tb.text_frame.paragraphs[0], spec["num"], 21, FB, DARK)
            tb = tbox(s, X0 + 1.55, y + 0.05, 22, 1.05, anchor="m")
            p = tb.text_frame.paragraphs[0]
            add_text(p, spec.get("subtitle", ""), 14, FB, DARK)
            if spec.get("subtitle_note"):
                add_text(p, "_" + spec["subtitle_note"], 10, FM, "404040")
            y += 1.3
        if spec.get("headline"):
            if spec.get("req") and not spec.get("num"):
                y += 0.75
            for k, c in enumerate((NAVY, SKY)):
                box(s, "OVAL", X0 + 0.1 + k * 0.28, y + 0.37, 0.3, 0.3, fill=c)
            tb = tbox(s, X0 + 0.85, y + 0.0, X1 - X0 - 0.9, 1.05, anchor="m")
            add_text(tb.text_frame.paragraphs[0], spec["headline"], 13.5, FB, DARK)
            y += 1.1
        top = max(y + 0.1, 3.5)
        bottom = CONTENT_BOTTOM
        if spec.get("conclusion"):
            bottom -= CONCL_H + 0.25
            cb = box(s, "RECTANGLE", X0, CONTENT_BOTTOM - CONCL_H, X1 - X0, CONCL_H, fill=NAVY)
            gradient(cb, "1A3A9C", "1E4FD6", 0)
            fill_text(cb, spec["conclusion"], 14, FB, "FFFFFF", em_blue="FFE699", em_red="FFE699")
        cols = spec.get("columns") or [{"w": 1, "blocks": spec.get("blocks", [])}]
        gap = 0.45
        total_w = X1 - X0 - gap * (len(cols) - 1)
        wsum = sum(c.get("w", 1) for c in cols)
        x = X0
        used = 0
        for c in cols:
            cw = total_w * c.get("w", 1) / wsum
            end = render_blocks(s, c.get("blocks", []), x, top, cw, bottom, spec)
            used = max(used, (end - top) / max(0.1, bottom - top))
            if c.get("arrow_after"):
                ar = box(s, "RIGHT_ARROW", x + cw + 0.02, (top + bottom) / 2 - 0.7, gap + 0.4, 1.4, fill=SKY)
            x += cw + gap
        self.fill.append((len(self.prs.slides), spec["title"], used))
        self.page_no(s)
        for ev in spec.get("eval", []):
            self.page_of.setdefault(ev, []).append(len(self.prs.slides))
        return s


def _ranges(pages):
    if not pages:
        return "-"
    out, a, b = [], pages[0], pages[0]
    for p in pages[1:] + [None]:
        if p is not None and p == b + 1:
            b = p
            continue
        out.append(f"{a}~{b}" if b > a else str(a))
        if p is not None:
            a = b = p
    return ", ".join(out)


# ---------------------------------------------------------------- 블록

def warn(spec, need, room, what):
    if need > room + 0.05:
        print(f"  ! 넘침 {need - room:.1f}cm [{spec.get('title')}] {what}", file=sys.stderr)


def render_blocks(s, blocks, x, y, w, bottom, spec):
    for i, b in enumerate(blocks):
        last = i == len(blocks) - 1
        kind = b.get("type", "bullets")
        room = bottom - y
        h = BLOCKS[kind](s, b, x, y, w, room if (last and b.get("fill", True)) else None, spec)
        y += h + 0.3
    return y - 0.3


def _bullets_height(items, size, w):
    h = 0.2
    for it in items:
        it = it if isinstance(it, dict) else {"text": it}
        h += est_lines(it["text"], size, w - 0.6) * lh(size, 1.2) + 0.18
        for sub in it.get("sub", []):
            h += est_lines(sub, size - 1, w - 1.0) * lh(size - 1, 1.2) + 0.08
    return h


def _draw_bullets(s, items, x, y, w, h, size, char="•", extra_pt=0):
    tb = tbox(s, x, y, w, h, inset=0.12)
    tf = tb.text_frame
    first = True
    for it in items:
        it = it if isinstance(it, dict) else {"text": it}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        pfmt(p, line=1.2, before=3 + extra_pt)
        bullet(p, it.get("char", char))
        add_text(p, it["text"], size)
        for sub in it.get("sub", []):
            sp = tf.add_paragraph()
            pfmt(sp, line=1.2)
            bullet(sp, "-", 0.75, 0.3, FM)
            add_text(sp, sub, size - 1, FM, "404040")


def b_bullets(s, b, x, y, w, room, spec):
    size = b.get("size", 11)
    h = _bullets_height(b["items"], size, w)
    if room is not None:
        warn(spec, h, room, "bullets")
        h = max(h, room) if b.get("frame") else h
    extra = 0
    if room is not None and b.get("spread", True) and room > h * 1.15:
        grow = min(1.25, (room / h) ** 0.5)
        size = size * grow
        h2 = _bullets_height(b["items"], size, w)
        extra = max(0, min(14, (room - h2) / len(b["items"]) / 0.0353 * 0.55))
        h = room
    if b.get("frame"):
        box(s, "RECTANGLE", x, y, w, h, line=LINE)
    _draw_bullets(s, b["items"], x, y, w, h, size, b.get("char", "•"), extra)
    return h


def b_panel(s, b, x, y, w, room, spec):
    """남색 제목 띠 + 테두리 상자. 안에 blocks 를 쌓는다."""
    title_h = 0.8
    inner = b.get("blocks", [])
    est = title_h + 0.3 + sum(_est_block(ib, w - 0.5) + 0.3 for ib in inner)
    h = est if room is None else max(est, room) if b.get("stretch", True) else est
    if room is not None:
        warn(spec, est, room, f"panel {b.get('title')}")
    box(s, "RECTANGLE", x, y, w, h, fill="FFFFFF", line=NAVY, lw=1.0)
    hd = box(s, "RECTANGLE", x, y, w, title_h, fill=b.get("color", NAVY))
    fill_text(hd, b["title"], 11, FB, "FFFFFF", align=b.get("title_align", "l"), inset=0.25)
    render_blocks(s, inner, x + 0.25, y + title_h + 0.25, w - 0.5, y + h - 0.2, spec)
    return h


def b_cards(s, b, x, y, w, room, spec):
    items = b["items"]
    cols = b.get("cols", min(3, len(items)))
    gap = 0.25
    cw = (w - gap * (cols - 1)) / cols
    size = b.get("size", 10)
    rows_n = -(-len(items) // cols)
    body_h = 0.3
    for it in items:
        d = it.get("desc", [])
        d = d if isinstance(d, list) else [d]
        body_h = max(body_h, sum(est_lines(t, size, cw - 0.7) * lh(size, 1.2) + 0.08 for t in d) + 0.25)
    ch = b.get("h") or (0.85 + body_h)
    if room is not None and b.get("fill", True) and b.get("grow", True):
        ch = max(ch, (room - gap * (rows_n - 1)) / rows_n)
    need = rows_n * ch + gap * (rows_n - 1)
    if room is not None:
        warn(spec, need, room, "cards")
    base_body = body_h
    if ch - 1.0 > base_body * 1.2:
        size = size * min(1.35, ((ch - 1.0) / base_body) ** 0.5)
    for i, it in enumerate(items):
        cx = x + (i % cols) * (cw + gap)
        cy = y + (i // cols) * (ch + gap)
        box(s, "ROUNDED_RECTANGLE", cx, cy, cw, ch, fill="FFFFFF", line=LINE, adj=0.05)
        tb = tbox(s, cx + 0.2, cy + 0.12, cw - 0.4, 0.7, anchor="m")
        add_text(tb.text_frame.paragraphs[0], it["title"], size + 1, FB, b.get("title_color", BLUE))
        d = it.get("desc", [])
        d = d if isinstance(d, list) else [d]
        if d:
            box(s, "RECTANGLE", cx + 0.2, cy + 0.82, cw - 0.4, 0.02, fill=LINE)
            tb = tbox(s, cx + 0.15, cy + 0.9, cw - 0.3, ch - 1.0, inset=0.08)
            gap_pt = min(10, max(2, ((ch - 1.0) - base_body * 1.3) / max(1, len(d)) / 0.0353 * 0.5))
            for k, t in enumerate(d):
                p = tb.text_frame.paragraphs[0] if k == 0 else tb.text_frame.add_paragraph()
                pfmt(p, line=1.2, after=gap_pt)
                bullet(p, "✓" if b.get("check", True) else "•", 0.4, 0.35, "Segoe UI Symbol")
                add_text(p, t, size)
    return need


def b_flow(s, b, x, y, w, room, spec):
    steps = b["steps"]
    n = len(steps)
    gap = 0.55
    sw = (w - gap * (n - 1)) / n
    size = b.get("size", 10)
    body_h = 0.3
    for st in steps:
        d = st.get("desc", [])
        d = d if isinstance(d, list) else [d]
        body_h = max(body_h, sum(est_lines(t, size, sw - 0.5) * lh(size, 1.2) + 0.06 for t in d) + 0.2)
    h = b.get("h") or (1.0 + body_h)
    if room is not None:
        warn(spec, h, room, "flow")
        if b.get("grow"):
            if room - 1.0 > body_h * 1.3:
                size = size * min(1.3, ((room - 1.0) / body_h) ** 0.5)
            h = room
    for i, st in enumerate(steps):
        sx = x + i * (sw + gap)
        box(s, "ROUNDED_RECTANGLE", sx, y, sw, h, fill=LIGHT, line="AFC4E6", adj=0.06)
        hd = box(s, "ROUNDED_RECTANGLE", sx, y, sw, 0.85, fill=b.get("head", NAVY), adj=0.15)
        fill_text(hd, st["title"], size + 0.5, FB, "FFFFFF")
        d = st.get("desc", [])
        d = d if isinstance(d, list) else [d]
        tb = tbox(s, sx + 0.12, y + 0.95, sw - 0.24, h - 1.0, inset=0.06)
        for k, t in enumerate(d):
            p = tb.text_frame.paragraphs[0] if k == 0 else tb.text_frame.add_paragraph()
            pfmt(p, "c" if len(d) == 1 else "l", 1.2, after=2)
            if len(d) > 1:
                bullet(p, "•", 0.3, 0.25)
            add_text(p, t, size)
        if i < n - 1:
            box(s, "RIGHT_ARROW", sx + sw + 0.08, y + h / 2 - 0.3, gap - 0.16, 0.6, fill=SKY)
    return h


def b_compare(s, b, x, y, w, room, spec):
    """As-Is / To-Be 비교: rows: [{label, asis:[...], tobe:[...]}]"""
    rows = b["rows"]
    lw_, gap = 2.6, 0.5
    cw = (w - lw_ - gap - 0.3) / 2
    size = b.get("size", 10)
    hy = y
    for k, (t, c) in enumerate(((b.get("asis_title", "As-Is"), "595959"), (b.get("tobe_title", "To-Be"), BLUE))):
        tb = tbox(s, x + lw_ + 0.3 + k * (cw + gap), hy, cw, 0.8, anchor="m")
        p = tb.text_frame.paragraphs[0]
        pfmt(p, "c")
        add_text(p, t, 12, FB, c)
    y0 = y + 0.9
    heights = []
    for r in rows:
        hh = max(_bullets_height(r["asis"], size, cw), _bullets_height(r["tobe"], size, cw)) + 0.15
        heights.append(max(1.6, hh))
    need = 0.9 + sum(heights) + 0.2 * (len(rows) - 1)
    if room is not None:
        warn(spec, need, room, "compare")
        if room > need and b.get("grow", True):
            add = (room - need) / len(rows)
            heights = [hh + add for hh in heights]
            need = room
    yy = y0
    for r, hh in zip(rows, heights):
        lb = box(s, "RECTANGLE", x, yy, lw_, hh, fill="FFFFFF", line=NAVY, lw=1.0)
        fill_text(lb, r["label"], 11, FB, DARK)
        box(s, "RECTANGLE", x + lw_ + 0.3, yy, cw, hh, fill="F3F4F6")
        _draw_bullets(s, r["asis"], x + lw_ + 0.3, yy + 0.05, cw, hh, size)
        box(s, "RIGHT_ARROW", x + lw_ + 0.3 + cw + 0.05, yy + hh / 2 - 0.3, gap - 0.1, 0.6, fill="9DB7E6")
        box(s, "RECTANGLE", x + lw_ + 0.3 + cw + gap, yy, cw, hh, fill="E3EEFC")
        _draw_bullets(s, r["tobe"], x + lw_ + 0.3 + cw + gap, yy + 0.05, cw, hh, size)
        yy += hh + 0.2
    return need


def b_layers(s, b, x, y, w, room, spec):
    layers = b["layers"]
    side_w = b.get("side_w", 5.0) if any(l.get("side") for l in layers) else 0
    main_w = w - (side_w + 0.25 if side_w else 0)
    n = len(layers)
    lh_ = b.get("h") or 1.35
    if room is not None and b.get("grow", True):
        lh_ = max(lh_, min(2.2, (room - 0.15 * (n - 1)) / n))
    need = n * lh_ + 0.15 * (n - 1)
    if room is not None:
        warn(spec, need, room, "layers")
    pal = ["1E3A8A", "2453B8", "1F6FD6", "2B86E0", "1C4E9E", "3A7BD5"]
    for i, l in enumerate(layers):
        ly = y + i * (lh_ + 0.15)
        col = l.get("color", pal[i % len(pal)])
        box(s, "RECTANGLE", x, ly, main_w, lh_, fill=col)
        nb = tbox(s, x + 0.1, ly, 3.2, lh_, anchor="m")
        p = nb.text_frame.paragraphs[0]
        pfmt(p, "c", 1.05)
        add_text(p, l["name"], 11, FB, "FFFFFF", em_blue="FFE699")
        items = l.get("items", [])
        if items:
            iw = (main_w - 3.5 - 0.15 * (len(items) - 1) - 0.15) / len(items)
            for k, it in enumerate(items):
                bx = box(s, "ROUNDED_RECTANGLE", x + 3.4 + k * (iw + 0.15), ly + 0.18, iw, lh_ - 0.36,
                         fill="FFFFFF", adj=0.08)
                fill_text(bx, it, 9.5, FB, "1F3864", inset=0.06)
        if l.get("side"):
            sb = box(s, "RECTANGLE", x + main_w + 0.25, ly, side_w, lh_, fill="F4F7FC", line=LINE)
            fill_text(sb, l["side"], 9.5, FM, "1F3864", inset=0.1)
    return need


def b_table(s, b, x, y, w, room, spec):
    rows = [b["header"]] + b["rows"]
    size = b.get("size", 9.5)
    widths = b.get("widths") or [1] * len(b["header"])
    k = w / sum(widths)
    widths = [c * k for c in widths]
    heights = []
    for i, r in enumerate(rows):
        lines = max(est_lines(c, size, widths[j] - 0.25) for j, c in enumerate(r))
        heights.append(max(0.75 if i else 0.7, lines * lh(size, 1.25) + 0.4))
    need = sum(heights)
    if room is not None and b.get("grow"):
        extra = (room - need) / len(rows)
        if extra > 0:
            heights = [hh + extra for hh in heights]
            need = room
    if room is not None:
        warn(spec, need, room, "table")
    table(s, x, y, w, rows, widths, size, heights, merge_first=b.get("merge_first", False),
          align=b.get("align"), first_bold=b.get("first_bold", True))
    return need


def table(s, x, y, w, rows, widths, size, heights, merge_first=False, header_fill=BLUE, align=None, first_bold=True):
    gf = s.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(w), Cm(sum(heights)))
    tbl = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Cm(cw)
    for i, hh in enumerate(heights):
        tbl.rows[i].height = Cm(hh)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Cm(0.12)
            cell.margin_top = cell.margin_bottom = Cm(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            for k, ln in enumerate(str(v).split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                left = i > 0 and (align == "l" or (align is None and j == len(r) - 1 and len(r) > 2)
                                  or (align is None and len(str(v)) > 28))
                pfmt(p, "l" if left else "c", 1.1)
                if i == 0:
                    add_text(p, ln, size + 0.5, FB, "FFFFFF", em_blue="FFE699")
                else:
                    add_text(p, ln, size, FB if (j == 0 and first_bold) else FM, "222222")
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(side))
                ln.set("w", "9525")
                sf = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", "BFC8D6")
            fill = header_fill if i == 0 else ("EEF3FB" if (j == 0 and first_bold) else "FFFFFF")
            sf = etree.SubElement(tcPr, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", fill)
    if merge_first:
        start = 1
        for i in range(2, len(rows) + 1):
            if i == len(rows) or rows[i][0] != rows[start][0]:
                if i - 1 > start and rows[start][0]:
                    for k in range(start + 1, i):
                        for r in tbl.cell(k, 0).text_frame.paragraphs[0].runs:
                            r.text = ""
                    tbl.cell(start, 0).merge(tbl.cell(i - 1, 0))
                start = i
    return gf


def b_kpi(s, b, x, y, w, room, spec):
    items = b["items"]
    n = len(items)
    gap = 0.3
    kw = (w - gap * (n - 1)) / n
    h = b.get("h", 2.6)
    for i, it in enumerate(items):
        kx = x + i * (kw + gap)
        box(s, "RECTANGLE", kx, y, kw, h, fill="FFFFFF", line=LINE)
        tb = tbox(s, kx, y + 0.15, kw, 1.3, anchor="m")
        p = tb.text_frame.paragraphs[0]
        pfmt(p, "c")
        add_text(p, it["value"], 24, FB, BLUE)
        tb = tbox(s, kx, y + 1.45, kw, h - 1.5, anchor="m")
        p = tb.text_frame.paragraphs[0]
        pfmt(p, "c")
        add_text(p, it["label"], 10, FB, DARK)
    return h


def b_text(s, b, x, y, w, room, spec):
    size = b.get("size", 11)
    h = est_lines(b["text"], size, w) * lh(size, 1.25) + 0.25
    tb = tbox(s, x, y, w, h)
    p = tb.text_frame.paragraphs[0]
    pfmt(p, b.get("align", "l"), 1.25)
    add_text(p, b["text"], size, FB if b.get("bold") else FM, b.get("color", DARK))
    return h


def b_label(s, b, x, y, w, room, spec):
    box(s, "RECTANGLE", x, y + 0.12, 0.1, 0.5, fill=DARK)
    tb = tbox(s, x + 0.2, y, w - 0.2, 0.75, anchor="m")
    add_text(tb.text_frame.paragraphs[0], b["text"], 12, FB, DARK)
    return 0.75


def b_callout(s, b, x, y, w, room, spec):
    size = b.get("size", 10.5)
    h = est_lines(b["text"], size, w - 0.6) * lh(size, 1.25) + 0.5
    bx = box(s, "ROUNDED_RECTANGLE", x, y, w, h, fill="FFF4E5", line="F4B183", adj=0.1)
    fill_text(bx, b["text"], size, FB, "7F3F00", align=b.get("align", "c"), inset=0.25, line=1.2)
    return h


def b_note(s, b, x, y, w, room, spec):
    tb = tbox(s, x, y, w, 0.55)
    add_text(tb.text_frame.paragraphs[0], b["text"], 9.5, FM, "666666")
    return 0.55


BLOCKS = {"bullets": b_bullets, "panel": b_panel, "cards": b_cards, "flow": b_flow, "compare": b_compare,
          "layers": b_layers, "table": b_table, "kpi": b_kpi, "text": b_text, "label": b_label,
          "callout": b_callout, "note": b_note}


def _est_block(b, w):
    """panel 안 블록 높이 추정 (그리지 않음)."""
    kind = b.get("type", "bullets")
    size = b.get("size", 10 if kind in ("cards", "flow") else 11)
    if kind == "bullets":
        return _bullets_height(b["items"], size, w)
    if kind == "table":
        widths = b.get("widths") or [1] * len(b["header"])
        k = w / sum(widths)
        tot = 0
        for i, r in enumerate([b["header"]] + b["rows"]):
            lines = max(est_lines(c, b.get("size", 9.5), widths[j] * k - 0.25) for j, c in enumerate(r))
            tot += max(0.75, lines * lh(b.get("size", 9.5), 1.2) + 0.3)
        return tot
    if kind == "cards":
        cols = b.get("cols", min(3, len(b["items"])))
        return (-(-len(b["items"]) // cols)) * (b.get("h") or 2.4)
    if kind == "flow":
        return b.get("h") or 2.4
    if kind == "layers":
        return len(b["layers"]) * ((b.get("h") or 1.35) + 0.15)
    if kind == "kpi":
        return b.get("h", 2.6)
    if kind in ("text", "callout"):
        return est_lines(b["text"], size, w) * lh(size, 1.25) + 0.5
    if kind == "label":
        return 0.75
    return 1.0


# ---------------------------------------------------------------- 조립

def build(path: Path, out: Path, report=False):
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    d = Deck(data)
    d.cover(data["meta"])
    if data.get("eval_items"):
        d.eval_table(data["eval_items"])
    d.contents(data["chapters"])
    for ch in data["chapters"]:
        d.divider(ch)
        for spec in ch.get("slides", []):
            kind = spec.get("kind", "body")
            if kind == "module":
                d.module(spec["title"])
            elif kind == "visual":
                d.visual(spec)
                for ev in spec.get("eval", []):
                    d.page_of.setdefault(ev, []).append(len(d.prs.slides))
            else:
                d.body(ch, spec)
    if data.get("closing"):
        d.closing(data["closing"])
    if data.get("eval_items"):
        d.fill_eval_table()
        missing = [i["item"] for i in data["eval_items"] if not d.page_of.get(i["item"]) and not i.get("page_text")]
        if missing:
            print("  ! 발표자료에 연결되지 않은 평가항목:", ", ".join(missing), file=sys.stderr)
    if report:
        for n, t, r in d.fill:
            print(f"  p{n:>3} {r * 100:5.0f}%  {t}")
    finalize(d.prs)
    set_pptx_author(d.prs)
    d.prs.save(str(out))
    print(f"저장: {out}  ({len(d.prs.slides)}장)")


def main(argv=None):
    ap = argparse.ArgumentParser(description="H9 제안 발표자료 PPT 생성")
    ap.add_argument("content", type=Path)
    ap.add_argument("-o", "--out", type=Path)
    ap.add_argument("--report", action="store_true")
    a = ap.parse_args(argv)
    build(a.content, a.out or a.content.with_suffix(".pptx"), a.report)


if __name__ == "__main__":
    main()
