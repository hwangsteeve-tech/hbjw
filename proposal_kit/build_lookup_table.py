"""제안요청서 대비 제안사항 조견표(엑셀) 생성기.

제안서 내용 파일(content.yaml)에서 슬라이드 쪽수를 계산하고, 조견표 설정 파일(lookup_spec.yaml)의
평가항목·RFP 요구사항과 연결해 RFP 예시 양식대로 엑셀을 만든다.

    python -m proposal_kit.build_lookup_table proposals/xxx/content.yaml proposals/xxx/lookup_spec.yaml -o 조견표.xlsx

쪽수 계산은 build_proposal.py 의 슬라이드 순서(표지 → 요구사항 분석표 → 목차 → 장별 간지·본문)와 같다.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from proposal_kit.company import load_content, set_xlsx_author

FONT = "맑은 고딕"
HEAD_FILL = PatternFill("solid", fgColor="FFF2CC")
ITEM_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def slide_pages(content: dict):
    """본문 슬라이드마다 (장 번호, 절 제목, h3, 요구사항 ID 목록, 쪽수)를 돌려준다."""
    page = 1  # 표지
    if content.get("req_table", True):
        n_ids = sum(len(c["ids"]) for c in content.get("req_categories", []))
        page += max(1, -(-n_ids // 44))
    page += 1  # 목차
    out = []
    by_chapter = content.get("page_numbering") == "chapter"
    for ch in content["chapters"]:
        page += 1  # 간지
        for k, spec in enumerate(ch.get("slides", []), start=1):
            page += 1
            ids = [r for r in re.split(r"[,\s]+", spec.get("req") or "") if r]
            out.append({"chapter": ch["num"], "chapter_title": ch.get("title", ""), "header": spec.get("header", ""),
                        "h3": spec.get("h3") or spec.get("h2", ""), "ids": ids,
                        "page": (ch["num"], k) if by_chapter else page})
    return out


def _pages_text(pages):
    if pages and isinstance(pages[0], tuple):  # 장별 쪽번호 (장, 쪽)
        out = []
        for ch in dict.fromkeys(c for c, _ in pages):
            nums = sorted({n for c, n in pages if c == ch})
            out.append(f"{ch}-" + _pages_text(nums).replace(", ", f", {ch}-").replace("~", f"~{ch}-"))
        return ", ".join(out)
    pages = sorted(set(pages))
    if not pages:
        return "-"
    # 연속 쪽은 범위로 묶는다 (예: 24~26)
    out, start, prev = [], pages[0], pages[0]
    for p in pages[1:] + [None]:
        if p is not None and p == prev + 1:
            prev = p
            continue
        out.append(f"{start}~{prev}" if prev > start else f"{start}")
        if p is not None:
            start = prev = p
    return ", ".join(out)


def _title(h3):
    return re.sub(r"^[\d.]+\s*", "", h3).strip()


def build_rows(content, spec):
    slides = slide_pages(content)
    reqs = spec.get("requirements") or {}
    rows = []
    for ev in spec["evaluation"]:
        item = f"{ev['item']}\n({ev['score']}점)" if ev["score"] != "-" else f"{ev['item']}\n(-)"
        for r in ev["rows"]:
            if "id" in r:
                rid = r["id"]
                q = reqs[rid]
                hit = [s for s in slides if rid in s["ids"]]
                rows.append([item, rid, str(q["page"]), f"{q['name']}\n- {q['summary']}",
                             _pages_text([s["page"] for s in hit]),
                             "\n".join(dict.fromkeys(_title(s["h3"]) for s in hit)) or "-", ""])
            else:
                h3s = r.get("h3") or []  # h3 번호 접두어로 특정 슬라이드만 지정 (예: ["2.1.1", "2.2.1"])
                hit = [s for s in slides if s["chapter"] == r["chapter"]
                       and (not r.get("header") or s["header"] == r["header"])
                       and (not h3s or any(s["h3"].startswith(str(h)) for h in h3s))]
                rows.append([item, r.get("label", "-"), r["rfp_page"], r["rfp_text"],
                             _pages_text([s["page"] for s in hit]),
                             "\n".join(dict.fromkeys(_title(s["h3"]) for s in hit)) or "-", ""])
    return rows


def build_compliance_rows(content, spec):
    """요구사항 고유번호 · 명칭 · 제안서 목차 · 쪽수 · 준수 여부 형식."""
    slides = slide_pages(content)
    rows = []
    for rid, q in spec["requirements"].items():
        hit = [s for s in slides if rid in s["ids"]]
        toc = "\n".join(dict.fromkeys(f"{s['chapter']}. {s['chapter_title']} > {s['h3']}" for s in hit)) or "-"
        rows.append([rid, q["name"], toc, _pages_text([s["page"] for s in hit]), q.get("comply", "준수"), q.get("note", "")])
    return rows


def write_compliance_xlsx(rows, spec, out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "조견표"
    head = spec.get("columns", ["요구사항 고유번호", "요구사항 명칭", "제안서 목차", "쪽수", "준수 여부", "비고"])
    widths = [14, 26, 52, 12, 10, 14]
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEF"[i]].width = w
    ws.merge_cells("A1:F1")
    ws["A1"] = spec.get("title", "제안요청서 요구사항 조견표")
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:F2")
    ws["A2"] = f"(제안사 : {spec.get('proposer', '')})"
    ws["A2"].font = Font(name=FONT, size=10)
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center")
    for j, h in enumerate(head):
        c = ws.cell(row=3, column=j + 1, value=h)
        c.font = Font(name=FONT, size=10, bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[3].height = 26
    for i, row in enumerate(rows):
        r = 4 + i
        for j, val in enumerate(row):
            c = ws.cell(row=r, column=j + 1, value=val)
            c.font = Font(name=FONT, size=9, bold=(j == 0))
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if j in (1, 2) else "center", vertical="center", wrap_text=True)
            if j == 0:
                c.fill = ITEM_FILL
        ws.row_dimensions[r].height = _height(row, widths)
    ws.freeze_panes = "A4"
    ws.print_title_rows = "3:3"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.oddFooter.center.text = "&P / &N"
    set_xlsx_author(wb)
    wb.save(out)


def _height(row, widths):
    lines = 1
    for val, w in zip(row, widths):
        text = str(val)
        cap = max(4, int(w * 1.05))  # 한글 기준 셀 폭(문자 수) 근사
        n = sum(max(1, -(-int(sum(1.0 if ord(c) > 0x2000 else 0.55 for c in seg)) // cap)) for seg in text.split("\n"))
        lines = max(lines, n)
    return max(20, lines * 14 + 6)


def write_xlsx(rows, spec, out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "조견표"
    widths = [13, 10, 9, 44, 9, 40, 8]
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEFG"[i]].width = w

    ws.merge_cells("A1:G1")
    ws["A1"] = spec.get("title", "제안요청서 대비 제안사항 조견표")
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:G2")
    ws["A2"] = f"(제안사 : {spec.get('proposer', '')})"
    ws["A2"].font = Font(name=FONT, size=10)
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center")

    head = [("A3:A4", "평가항목\n(배점)"), ("B3:B4", "요구사항\n구분"), ("C3:D3", "제안요청서 요구사항"),
            ("C4", spec.get("rfp_ref_header", "쪽수")), ("D4", "요구사항명 및\n주요내용"), ("E3:F3", "제안사항"), ("E4", "쪽수"),
            ("F4", "주요내용"), ("G3:G4", "비고")]
    for rng, text in head:
        if ":" in rng:
            ws.merge_cells(rng)
        cell = ws[rng.split(":")[0]]
        cell.value = text
    for r in (3, 4):
        ws.row_dimensions[r].height = 24
        for c in "ABCDEFG":
            cell = ws[f"{c}{r}"]
            cell.font = Font(name=FONT, size=10, bold=True)
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BOX

    start = 5
    for i, row in enumerate(rows):
        r = start + i
        for j, val in enumerate(row):
            cell = ws.cell(row=r, column=j + 1, value=val)
            cell.font = Font(name=FONT, size=9, bold=(j == 0))
            cell.border = BOX
            left = j in (3, 5)
            cell.alignment = Alignment(horizontal="left" if left else "center", vertical="center", wrap_text=True)
            if j == 0:
                cell.fill = ITEM_FILL
        ws.row_dimensions[r].height = _height(row, widths)

    # 같은 평가항목은 첫 열을 병합
    r0 = start
    for i in range(1, len(rows) + 1):
        if i == len(rows) or rows[i][0] != rows[r0 - start][0]:
            if start + i - 1 > r0:
                ws.merge_cells(start_row=r0, start_column=1, end_row=start + i - 1, end_column=1)
            r0 = start + i

    ws.freeze_panes = "A5"
    ws.print_title_rows = "3:4"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.oddFooter.center.text = "&P / &N"
    set_xlsx_author(wb)
    wb.save(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description="제안요청서 대비 제안사항 조견표 생성")
    ap.add_argument("content", type=Path)
    ap.add_argument("spec", type=Path)
    ap.add_argument("-o", "--out", type=Path, required=True)
    args = ap.parse_args(argv)
    content = load_content(args.content)
    spec = yaml.safe_load(args.spec.read_text(encoding="utf-8"))
    if spec.get("format") == "compliance":
        rows = build_compliance_rows(content, spec)
        missing = [r[0] for r in rows if r[3] == "-"]
        write = write_compliance_xlsx
    else:
        rows = build_rows(content, spec)
        missing = [r[1] for r in rows if r[4] == "-"]
        write = write_xlsx
    if missing:
        print("  ! 제안서 쪽수가 없는 항목:", ", ".join(missing))
    write(rows, spec, args.out)
    print(f"저장: {args.out}  ({len(rows)}행)")


if __name__ == "__main__":
    main()
