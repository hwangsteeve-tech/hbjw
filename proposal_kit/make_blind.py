"""업체명 미표기(블라인드) 제출본을 만든다.

빌드한 정성제안서 · 발표자료 PPT에서 회사를 알아볼 수 있는 글자(회사명 · 영문명 · 대표자 · 주소 · 연락처 ·
사업자등록번호)를 지우고 '제안사' / '○○○'로 바꾼다. 표 · 그룹 안의 글자도 바꾸며, 레이아웃 · 글꼴은 그대로 둔다.
바꾼 뒤에도 회사를 알아볼 수 있는 글자가 남아 있으면 쪽 번호와 함께 알려 준다.

    python -m proposal_kit.make_blind proposals/xxx/xxx_정성제안서.pptx -o proposals/xxx/블라인드_정성제안서.pptx
    python -m proposal_kit.make_blind 입력.pptx -o 출력.pptx --term "추가로 지울 문구=바꿀 문구"
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation

from proposal_kit.company import load_company, set_pptx_author


def blind_terms(extra: list[str] | None = None) -> list[tuple[str, str]]:
    """(지울 문구, 바꿀 문구) — 긴 문구부터 바꾸도록 정렬한다."""
    g = load_company()["general"]
    name = g["name"]                              # 주식회사 에이치나인
    short = name.replace("주식회사", "").strip()   # 에이치나인
    pairs = [
        (f"제출기관: {name}", "제출기관: 제안사"),
        (f"{name} ({g['name_en']})", "제안사"),
        (name, "제안사"),
        (f"㈜{short}", "제안사"),
        (f"(주){short}", "제안사"),
        (f"{short}(HNINE)", "제안사"),
        (short, "제안사"),
        (g["name_en"], "제안사"),
        (g["ceo"], "○○○"),
        (g["address"], "○○○"),
        (g["tel"], "○○○"),
        (g["fax"], "○○○"),
        (g["biz_no"], "○○○"),
        ("H9", "제안사"),
    ]
    for t in extra or []:
        old, _, new = t.partition("=")
        pairs.append((old, new or "○○○"))
    return sorted(pairs, key=lambda p: -len(p[0]))


def _paragraphs(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # 그룹
            yield from _paragraphs(sh.shapes)
            continue
        if sh.has_text_frame:
            yield from sh.text_frame.paragraphs
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs


def _replace_paragraph(p, pairs) -> int:
    """문단 단위로 바꾼다. 문구가 여러 run에 걸쳐 있으면 첫 run에 합쳐 서식을 유지한다."""
    runs = p.runs
    if not runs:
        return 0
    text = "".join(r.text for r in runs)
    new = text
    for old, rep in pairs:
        if old == "H9":
            new = re.sub(r"(?<![A-Za-z0-9])H9(?![A-Za-z0-9])", rep, new)
        else:
            new = new.replace(old, rep)
    if new == text:
        return 0
    # run 경계를 최대한 살린다: run별로 먼저 바꿔 보고, 안 되면 첫 run에 몰아 넣는다
    for r in runs:
        t = r.text
        for old, rep in pairs:
            t = re.sub(r"(?<![A-Za-z0-9])H9(?![A-Za-z0-9])", rep, t) if old == "H9" else t.replace(old, rep)
        r.text = t
    if "".join(r.text for r in runs) != new:
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""
    return 1


def make_blind(src: Path, out: Path, extra: list[str] | None = None) -> int:
    pairs = blind_terms(extra)
    prs = Presentation(str(src))
    n = 0
    for s in prs.slides:
        for p in _paragraphs(s.shapes):
            n += _replace_paragraph(p, pairs)
        if s.has_notes_slide:
            for p in s.notes_slide.notes_text_frame.paragraphs:
                n += _replace_paragraph(p, pairs)
    # 남은 식별 문구 점검
    left = []
    checks = [o for o, _ in pairs if o not in ("H9",)] + ["HNINE", "에이치나인"]
    for i, s in enumerate(prs.slides, start=1):
        txt = "\n".join(p.text for p in _paragraphs(s.shapes))
        hit = [c for c in checks if c and c in txt]
        if re.search(r"(?<![A-Za-z0-9])H9(?![A-Za-z0-9])", txt):
            hit.append("H9")
        if hit:
            left.append((i, hit))
    for i, hit in left:
        print(f"  ! {i}쪽에 남은 식별 문구: {', '.join(hit)}", file=sys.stderr)
    set_pptx_author(prs)
    prs.save(str(out))
    print(f"저장: {out}  (문단 {n}곳 바꿈{', 남은 문구 없음' if not left else ''})")
    return len(left)


def main(argv=None):
    ap = argparse.ArgumentParser(description="업체명 미표기(블라인드) 제출본 생성")
    ap.add_argument("src", type=Path)
    ap.add_argument("-o", "--out", type=Path, required=True)
    ap.add_argument("--term", action="append", help="추가로 지울 문구 (예: '홍길동=○○○')")
    args = ap.parse_args(argv)
    sys.exit(1 if make_blind(args.src, args.out, args.term) else 0)


if __name__ == "__main__":
    main()
